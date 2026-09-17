import os
from datetime import datetime
from io import BytesIO

import streamlit as st
from dotenv import load_dotenv
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from database import (
    clear_conversation,
    create_conversation,
    create_tables,
    delete_conversation,
    get_conversation,
    get_conversations,
    get_messages,
    remove_document,
    save_document,
    save_message,
    update_conversation_title,
)


# ==================================================
# PAGE CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="Nova AI",
    page_icon="✨",
    layout="wide",
    initial_sidebar_state="expanded",
)

load_dotenv()
create_tables()


# ==================================================
# STYLING
# ==================================================

st.markdown(
    """
<style>
@import url(
    'https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap'
);

html,
body,
[class*="css"] {
    font-family: "Inter", sans-serif;
}

.stApp {
    background: #fbfbfd;
}

.block-container {
    max-width: 1050px;
    padding-top: 1.2rem;
    padding-bottom: 8rem;
}

[data-testid="stHeader"] {
    background: transparent;
}

#MainMenu {
    visibility: hidden;
}

footer {
    visibility: hidden;
}

/* Sidebar */

[data-testid="stSidebar"] {
    background: linear-gradient(
        180deg,
        #070b2b 0%,
        #090d32 55%,
        #0b1038 100%
    );
    border-right: 1px solid rgba(255, 255, 255, 0.05);
}

[data-testid="stSidebar"] > div:first-child {
    padding-top: 1.2rem;
}

[data-testid="stSidebar"] * {
    color: #f8fafc;
}

[data-testid="stSidebar"] hr {
    border-color: rgba(255, 255, 255, 0.08);
}

[data-testid="stSidebar"] .stCaption {
    color: #8f96ad !important;
}

[data-testid="stSidebar"] button {
    min-height: 42px;
    border-radius: 11px;
    border: 1px solid rgba(255, 255, 255, 0.07);
    background: transparent;
    color: #e6e9f2;
    text-align: left;
    transition: all 0.2s ease;
}

[data-testid="stSidebar"] button:hover {
    background: rgba(255, 255, 255, 0.08);
    border-color: rgba(139, 92, 246, 0.45);
    transform: translateY(-1px);
}

[data-testid="stSidebar"] button[kind="primary"] {
    background: linear-gradient(
        100deg,
        #1689ee 0%,
        #8753f6 100%
    );
    border: none;
    color: white;
    box-shadow: 0 8px 20px rgba(92, 82, 240, 0.25);
}

[data-testid="stSidebar"] input {
    background: #14193f !important;
    color: white !important;
    border: 1px solid #292f59 !important;
    border-radius: 11px !important;
}

[data-testid="stSidebar"] input::placeholder {
    color: #7f879f;
}

/* Chat messages */

[data-testid="stChatMessage"] {
    background: white;
    border: 1px solid #ececf2;
    border-radius: 18px;
    padding: 1rem 1.15rem;
    margin-bottom: 0.9rem;
    box-shadow: 0 4px 18px rgba(25, 25, 55, 0.035);
}

/* Input */

[data-testid="stChatInput"] {
    background: white;
    border: 1px solid #dddde7;
    border-radius: 20px;
    box-shadow: 0 10px 35px rgba(38, 37, 70, 0.10);
}

[data-testid="stChatInput"]:focus-within {
    border-color: #8b5cf6;
    box-shadow: 0 10px 35px rgba(124, 58, 237, 0.14);
}

/* Header */

.nova-header {
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding-bottom: 15px;
    border-bottom: 1px solid #eeeef3;
    margin-bottom: 25px;
}

.nova-header-title {
    color: #161622;
    font-size: 0.95rem;
    font-weight: 600;
    margin: 0;
}

.nova-header-status {
    color: #9295a5;
    font-size: 0.75rem;
    margin-top: 3px;
}

.ready-badge {
    display: inline-flex;
    align-items: center;
    gap: 7px;
    padding: 7px 12px;
    border: 1px solid #e7e7ef;
    border-radius: 20px;
    background: white;
    color: #606474;
    font-size: 0.76rem;
}

.ready-dot {
    width: 7px;
    height: 7px;
    background: #22c55e;
    border-radius: 50%;
    display: inline-block;
}

/* Welcome */

.welcome-section {
    text-align: center;
    max-width: 720px;
    margin: 4rem auto 2rem auto;
}

.nova-logo {
    width: 49px;
    height: 49px;
    margin: auto;
    border-radius: 15px;
    display: flex;
    align-items: center;
    justify-content: center;
    color: white;
    font-size: 24px;
    background: linear-gradient(
        135deg,
        #208fee 0%,
        #8654f6 100%
    );
    box-shadow: 0 10px 25px rgba(109, 82, 235, 0.25);
}

.meet-nova {
    color: #8757e8;
    font-size: 0.88rem;
    font-weight: 600;
    margin-top: 18px;
    margin-bottom: 7px;
}

.welcome-title {
    color: #171720;
    font-size: 2.45rem;
    font-weight: 750;
    letter-spacing: -0.045em;
    margin: 0;
}

.welcome-description {
    max-width: 620px;
    margin: 14px auto 0 auto;
    color: #777987;
    font-size: 1rem;
    line-height: 1.7;
}

div[data-testid="stHorizontalBlock"] button {
    border-radius: 15px;
}

/* Sidebar components */

.pdf-card {
    padding: 13px 14px;
    border-radius: 12px;
    background: #151a43;
    border: 1px solid #252b58;
    margin-bottom: 10px;
    overflow-wrap: anywhere;
}

.pdf-card-label {
    color: #7e87a4;
    font-size: 0.68rem;
    letter-spacing: 0.08em;
    text-transform: uppercase;
    margin-bottom: 7px;
}

.pdf-card-name {
    color: #f2f3f8;
    font-size: 0.82rem;
}

.section-label {
    color: #7f879d;
    font-size: 0.68rem;
    font-weight: 600;
    letter-spacing: 0.09em;
    text-transform: uppercase;
    margin: 12px 0 9px 5px;
}

.brand-row {
    display: flex;
    align-items: center;
    gap: 11px;
    margin-bottom: 20px;
}

.brand-icon {
    width: 37px;
    height: 37px;
    border-radius: 11px;
    display: flex;
    align-items: center;
    justify-content: center;
    background: linear-gradient(
        135deg,
        #238eea 0%,
        #8753f5 100%
    );
    font-size: 18px;
    box-shadow: 0 7px 18px rgba(83, 82, 235, 0.28);
}

.brand-name {
    font-size: 1.06rem;
    font-weight: 700;
    color: white;
}

.disclaimer {
    text-align: center;
    color: #999ca9;
    font-size: 0.7rem;
    margin-top: 9px;
}

.supported-files {
    color: #8f96ad;
    font-size: 0.72rem;
    line-height: 1.5;
    margin-top: 8px;
}

/* Mobile */

@media (max-width: 768px) {
    .block-container {
        padding-left: 1rem;
        padding-right: 1rem;
    }

    .welcome-section {
        margin-top: 2.5rem;
    }

    .welcome-title {
        font-size: 1.9rem;
    }

    .welcome-description {
        font-size: 0.9rem;
    }
}
/* Hide Streamlit hover tooltips */
[data-testid="stTooltipContent"],
div[role="tooltip"],
[data-baseweb="tooltip"] {
    display: none !important;
    visibility: hidden !important;
    opacity: 0 !important;
    pointer-events: none !important;
}
</style>
    """,
    unsafe_allow_html=True,
)


# ==================================================
# GEMINI CLIENT
# ==================================================

api_key = os.getenv("GEMINI_API_KEY")

if api_key:
    client = genai.Client(api_key=api_key)
else:
    client = None

model_name = os.getenv(
    "GEMINI_MODEL",
    "gemini-3.1-flash-lite",
)


# ==================================================
# FILE EXTRACTION FUNCTIONS
# ==================================================

def extract_pdf_text(file_bytes):
    reader = PdfReader(BytesIO(file_bytes))
    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        page_text = page.extract_text()

        if page_text:
            pages.append(
                f"--- PDF Page {page_number} ---\n"
                f"{page_text}"
            )

    return "\n\n".join(pages).strip()


def extract_word_text(file_bytes):
    document = Document(BytesIO(file_bytes))
    content = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            content.append(paragraph_text)

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):
        content.append(
            f"\n--- Word Table {table_number} ---"
        )

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            content.append(" | ".join(values))

    return "\n".join(content).strip()


def extract_excel_text(file_bytes):
    workbook = load_workbook(
        BytesIO(file_bytes),
        data_only=True,
        read_only=True,
    )

    content = []

    for worksheet in workbook.worksheets:
        content.append(
            f"\n--- Excel Sheet: {worksheet.title} ---"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = [
                "" if value is None else str(value)
                for value in row
            ]

            if any(value.strip() for value in values):
                content.append(" | ".join(values))

    workbook.close()

    return "\n".join(content).strip()


def extract_powerpoint_text(file_bytes):
    presentation = Presentation(BytesIO(file_bytes))
    content = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        content.append(
            f"\n--- PowerPoint Slide {slide_number} ---"
        )

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                shape_text = shape.text.strip()

                if shape_text:
                    content.append(shape_text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    content.append(" | ".join(values))

    return "\n".join(content).strip()


def analyze_image(file_bytes, mime_type):
    if not client:
        raise ValueError(
            "Gemini API key is required to analyze images."
        )

    if len(file_bytes) > 15 * 1024 * 1024:
        raise ValueError(
            "The image is too large. Use an image "
            "smaller than 15 MB."
        )

    image_part = types.Part.from_bytes(
        data=file_bytes,
        mime_type=mime_type,
    )

    response = client.models.generate_content(
        model=model_name,
        contents=[
            (
                "Analyze this image carefully. Describe the "
                "important content, including visible text, "
                "objects, people, tables, diagrams, charts, "
                "labels and context. This description will be "
                "used to answer future questions about the image."
            ),
            image_part,
        ],
    )

    return (
        response.text
        or "No image description was generated."
    )


def extract_uploaded_file(uploaded_file):
    file_name = uploaded_file.name
    extension = file_name.rsplit(".", 1)[-1].lower()
    file_bytes = uploaded_file.getvalue()

    if extension == "pdf":
        extracted_content = extract_pdf_text(
            file_bytes
        )

    elif extension == "docx":
        extracted_content = extract_word_text(
            file_bytes
        )

    elif extension == "xlsx":
        extracted_content = extract_excel_text(
            file_bytes
        )

    elif extension == "pptx":
        extracted_content = extract_powerpoint_text(
            file_bytes
        )

    elif extension in [
        "png",
        "jpg",
        "jpeg",
        "webp",
    ]:
        mime_type = uploaded_file.type

        if not mime_type:
            if extension in ["jpg", "jpeg"]:
                mime_type = "image/jpeg"
            else:
                mime_type = f"image/{extension}"

        extracted_content = analyze_image(
            file_bytes,
            mime_type,
        )

    else:
        raise ValueError(
            f"Unsupported file type: .{extension}"
        )

    if not extracted_content.strip():
        raise ValueError(
            "No readable content was found in this file."
        )

    return extracted_content


# ==================================================
# CHAT HELPER FUNCTIONS
# ==================================================

def make_conversation_title(message):
    cleaned_message = " ".join(message.split())

    if not cleaned_message:
        return "New conversation"

    if len(cleaned_message) > 30:
        return cleaned_message[:30] + "..."

    return cleaned_message


def format_conversation_date(created_at):
    if not created_at:
        return ""

    try:
        date_value = datetime.fromisoformat(
            str(created_at)
        )

        today = datetime.now().date()
        message_date = date_value.date()

        if message_date == today:
            return "Today"

        difference = (today - message_date).days

        if difference == 1:
            return "Yesterday"

        return date_value.strftime("%b %d")

    except (ValueError, TypeError):
        return ""


def load_conversation(conversation_id):
    conversation = get_conversation(conversation_id)
    saved_messages = get_messages(conversation_id)

    st.session_state.conversation_id = (
        conversation_id
    )

    st.session_state.messages = []

    for message in saved_messages:
        st.session_state.messages.append(
            {
                "role": message["role"],
                "content": message["content"],
                "document_name": message[
                    "document_name"
                ],
            }
        )

    if conversation:
        st.session_state.document_name = (
            conversation["document_name"]
        )

        st.session_state.document_text = (
            conversation["document_text"]
        )

    else:
        st.session_state.document_name = None
        st.session_state.document_text = None


def start_new_conversation():
    conversation_id = create_conversation()
    load_conversation(conversation_id)


def build_prompt(user_message):
    sections = [
        (
            "You are Nova, a helpful, friendly and "
            "professional AI assistant. Give clear, accurate "
            "and well-structured answers."
        )
    ]

    # Exclude the message just added, preventing duplication
    previous_messages = (
        st.session_state.messages[:-1][-10:]
    )

    if previous_messages:
        sections.append("\nConversation history:")

        for message in previous_messages:
            if message["role"] == "user":
                speaker = "User"
            else:
                speaker = "Assistant"

            sections.append(
                f"{speaker}: {message['content']}"
            )

    if st.session_state.document_text:
        attachment_content = (
            st.session_state.document_text[:50000]
        )

        sections.append(
            "\nCurrent attachment content:"
        )

        sections.append(attachment_content)

        sections.append(
            "\nUse the attachment when it is relevant "
            "to the user's question."
        )

    sections.append(
        f"\nCurrent user message: {user_message}"
    )

    return "\n".join(sections)


def escape_html(value):
    if not value:
        return ""

    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


# ==================================================
# SESSION STATE
# ==================================================

if "conversation_id" not in st.session_state:
    existing_conversations = get_conversations()

    if existing_conversations:
        first_conversation_id = (
            existing_conversations[0]["id"]
        )
    else:
        first_conversation_id = (
            create_conversation()
        )

    load_conversation(first_conversation_id)

if "messages" not in st.session_state:
    load_conversation(
        st.session_state.conversation_id
    )

if "document_name" not in st.session_state:
    st.session_state.document_name = None

if "document_text" not in st.session_state:
    st.session_state.document_text = None


# ==================================================
# SIDEBAR
# ==================================================

with st.sidebar:
    brand_html = (
        '<div class="brand-row">'
        '<div class="brand-icon">✣</div>'
        '<div class="brand-name">Nova AI</div>'
        '</div>'
    )

    st.markdown(
        brand_html,
        unsafe_allow_html=True,
    )

    if st.button(
        "＋  New chat",
        use_container_width=True,
        type="primary",
    ):
        start_new_conversation()
        st.rerun()

    search_text = st.text_input(
        "Search conversations",
        placeholder="Search conversations",
        label_visibility="collapsed",
    )

    st.markdown(
        '<div class="section-label">'
        'Recent conversations'
        '</div>',
        unsafe_allow_html=True,
    )

    conversations = get_conversations()

    if search_text:
        conversations = [
            conversation
            for conversation in conversations
            if search_text.lower()
            in conversation["title"].lower()
        ]

    if not conversations:
        st.caption("No conversations found")

    for conversation in conversations:
        conversation_id = conversation["id"]
        conversation_title = conversation["title"]

        conversation_date = format_conversation_date(
            conversation["created_at"]
        )

        selected = (
            conversation_id
            == st.session_state.conversation_id
        )

        chat_column, delete_column = st.columns(
            [5.3, 1],
            vertical_alignment="center",
        )

        with chat_column:
            chat_label = (
                f"💬  {conversation_title}"
            )

            if conversation_date:
                chat_label += (
                    f" · {conversation_date}"
                )

            if st.button(
                chat_label,
                key=f"open_chat_{conversation_id}",
                use_container_width=True,
                type=(
                    "primary"
                    if selected
                    else "secondary"
                ),
            ):
                load_conversation(
                    conversation_id
                )

                st.rerun()

        with delete_column:
            if st.button(
                "×",
                key=f"delete_chat_{conversation_id}",
                help=f"Delete {conversation_title}",
                use_container_width=True,
            ):
                deleting_current_chat = (
                    conversation_id
                    == st.session_state.conversation_id
                )

                delete_conversation(
                    conversation_id
                )

                if deleting_current_chat:
                    remaining = get_conversations()

                    if remaining:
                        next_id = remaining[0]["id"]
                    else:
                        next_id = create_conversation()

                    load_conversation(next_id)

                st.rerun()

    st.divider()

    st.markdown(
        '<div class="section-label">'
        'Current attachment'
        '</div>',
        unsafe_allow_html=True,
    )

    if st.session_state.document_name:
        safe_document_name = escape_html(
            st.session_state.document_name
        )

        attachment_html = (
            '<div class="pdf-card">'
            '<div class="pdf-card-label">'
            'Attached file'
            '</div>'
            '<div class="pdf-card-name">'
            f'📎 {safe_document_name}'
            '</div>'
            '</div>'
        )

        st.markdown(
            attachment_html,
            unsafe_allow_html=True,
        )

        if st.button(
            "×  Remove attachment",
            use_container_width=True,
        ):
            remove_document(
                st.session_state.conversation_id
            )

            st.session_state.document_name = None
            st.session_state.document_text = None

            st.rerun()

    else:
        st.caption(
            "Attach a document or image using "
            "the message box."
        )

    st.markdown(
        '<div class="supported-files">'
        'Supported: PDF, DOCX, XLSX, PPTX, '
        'PNG, JPG, JPEG and WEBP'
        '</div>',
        unsafe_allow_html=True,
    )

    if st.button(
        "🗑  Clear current conversation",
        use_container_width=True,
    ):
        clear_conversation(
            st.session_state.conversation_id
        )

        load_conversation(
            st.session_state.conversation_id
        )

        st.rerun()

    st.markdown(
        "<br>",
        unsafe_allow_html=True,
    )

    st.caption("👤  Nasra")
    st.caption("Personal workspace")


# ==================================================
# CONVERSATION HEADER
# ==================================================

current_conversation = get_conversation(
    st.session_state.conversation_id
)

if current_conversation:
    current_title = current_conversation["title"]
else:
    current_title = "New conversation"

safe_current_title = escape_html(current_title)

header_html = (
    '<div class="nova-header">'
    '<div>'
    f'<p class="nova-header-title">{safe_current_title}</p>'
    '<div class="nova-header-status">'
    'Nova 2.0 · Personal assistant'
    '</div>'
    '</div>'
    '<div class="ready-badge">'
    '<span class="ready-dot"></span>'
    'Ready'
    '</div>'
    '</div>'
)

st.markdown(
    header_html,
    unsafe_allow_html=True,
)


# ==================================================
# WELCOME SCREEN
# ==================================================

suggested_prompt = None

if not st.session_state.messages:
    welcome_html = (
        '<div class="welcome-section">'
        '<div class="nova-logo">✣</div>'
        '<div class="meet-nova">Meet Nova</div>'
        '<h1 class="welcome-title">'
        'How can I help you today?'
        '</h1>'
        '<p class="welcome-description">'
        'Ask a question, explore an idea, or attach '
        'a document or image. I’m here to help you '
        'think and create.'
        '</p>'
        '</div>'
    )

    st.markdown(
        welcome_html,
        unsafe_allow_html=True,
    )

    first_column, second_column = st.columns(2)

    with first_column:
        if st.button(
            "💡 Brainstorm ideas\n\n"
            "Plan a creative product launch",
            use_container_width=True,
        ):
            suggested_prompt = (
                "Help me brainstorm ideas for a "
                "creative product launch."
            )

        if st.button(
            "✨ Improve my writing\n\n"
            "Make my writing clearer",
            use_container_width=True,
        ):
            suggested_prompt = (
                "Help me improve my writing and make "
                "it clearer and more confident."
            )

    with second_column:
        if st.button(
            "📄 Summarize a document\n\n"
            "Pull out the important points",
            use_container_width=True,
        ):
            if st.session_state.document_text:
                suggested_prompt = (
                    "Summarize the attached file and "
                    "list its most important points."
                )
            else:
                st.warning(
                    "Please attach a document using "
                    "the message box."
                )

        if st.button(
            "💬 Learn something\n\n"
            "Explain a complex topic simply",
            use_container_width=True,
        ):
            suggested_prompt = (
                "Explain artificial intelligence "
                "in simple terms."
            )


# ==================================================
# DISPLAY SAVED MESSAGES
# ==================================================

for message in st.session_state.messages:
    if message["role"] == "user":
        avatar = "👤"
    else:
        avatar = "✨"

    with st.chat_message(
        message["role"],
        avatar=avatar,
    ):
        st.markdown(message["content"])

        if (
            message["role"] == "user"
            and message.get("document_name")
        ):
            st.caption(
                f"📎 {message['document_name']}"
            )


# ==================================================
# MESSAGE INPUT
# ==================================================

submission = st.chat_input(
    "Message Nova or attach a file...",
    accept_file=True,
    file_type=[
        "pdf",
        "docx",
        "xlsx",
        "pptx",
        "png",
        "jpg",
        "jpeg",
        "webp",
    ],
)

user_message = ""
uploaded_files = []

if suggested_prompt:
    user_message = suggested_prompt

elif submission:
    user_message = submission.text.strip()
    uploaded_files = submission.files


# ==================================================
# PROCESS THE MESSAGE AND FILE
# ==================================================

if user_message or uploaded_files:
    uploaded_document_name = None

    if uploaded_files:
        uploaded_file = uploaded_files[0]

        try:
            with st.spinner(
                f"Reading {uploaded_file.name}..."
            ):
                extracted_content = (
                    extract_uploaded_file(
                        uploaded_file
                    )
                )

            uploaded_document_name = (
                uploaded_file.name
            )

            st.session_state.document_name = (
                uploaded_document_name
            )

            st.session_state.document_text = (
                extracted_content
            )

            save_document(
                st.session_state.conversation_id,
                uploaded_document_name,
                extracted_content,
            )

        except Exception as error:
            st.error(
                f"Could not read the file: {error}"
            )

            st.stop()

    if not user_message and uploaded_document_name:
        file_extension = (
            uploaded_document_name
            .rsplit(".", 1)[-1]
            .lower()
        )

        if file_extension in [
            "png",
            "jpg",
            "jpeg",
            "webp",
        ]:
            user_message = (
                "Please explain what is shown "
                "in this image."
            )
        else:
            user_message = (
                "Please summarize the "
                "attached document."
            )

    if not user_message:
        st.stop()

    # Create conversation title from first message
    if not st.session_state.messages:
        new_title = make_conversation_title(
            user_message
        )

        update_conversation_title(
            st.session_state.conversation_id,
            new_title,
        )

    # Save user message
    user_data = {
        "role": "user",
        "content": user_message,
        "document_name": uploaded_document_name,
    }

    st.session_state.messages.append(user_data)

    save_message(
        st.session_state.conversation_id,
        "user",
        user_message,
        uploaded_document_name,
    )

    with st.chat_message(
        "user",
        avatar="👤",
    ):
        st.markdown(user_message)

        if uploaded_document_name:
            st.caption(
                f"📎 {uploaded_document_name}"
            )

    # Generate AI response
    with st.chat_message(
        "assistant",
        avatar="✨",
    ):
        if not client:
            assistant_message = (
                "GEMINI_API_KEY was not found. "
                "Please add it to your .env file."
            )

            st.error(assistant_message)

        else:
            try:
                with st.spinner(
                    "Nova is thinking..."
                ):
                    prompt = build_prompt(
                        user_message
                    )

                    response = (
                        client.models.generate_content(
                            model=model_name,
                            contents=prompt,
                        )
                    )

                    assistant_message = (
                        response.text
                        or (
                            "I could not generate "
                            "a response."
                        )
                    )

                st.markdown(assistant_message)

            except Exception as error:
                assistant_message = (
                    f"Something went wrong: {error}"
                )

                st.error(assistant_message)

    # Save assistant response
    assistant_data = {
        "role": "assistant",
        "content": assistant_message,
        "document_name": None,
    }

    st.session_state.messages.append(
        assistant_data
    )

    save_message(
        st.session_state.conversation_id,
        "assistant",
        assistant_message,
    )

    st.rerun()


# ==================================================
# DISCLAIMER
# ==================================================

st.markdown(
    '<div class="disclaimer">'
    'Nova can make mistakes. '
    'Check important information.'
    '</div>',
    unsafe_allow_html=True,
)